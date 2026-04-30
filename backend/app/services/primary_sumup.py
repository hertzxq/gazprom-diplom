"""
Первичный свод (Задача 2, шаг 1).

Агрегирует платежи по договорам, подтягивает реквизиты из реестра договоров,
классифицирует «Исключение из СМСП» и генерирует xlsx с 15 колонками A–O.

Колонки первичного свода:
    A — Дата (первая дата платежа в группе)
    B — Контрагент
    C — Номер договора
    D — Дата договора
    E — Сумма по договору
    F — Сумма платежа (сумма всех платежей по договору, в рублях)
    G — Вид СМСП                        (из реестра договоров, столбец G)
    H — Закупка для СМСП                (из реестра договоров, столбец H)
    I — Способ закупки                  (из реестра договоров, столбец F: ЕП/КИМ/МИ/…)
    J — Исключение из СМСП              (короткая метка — результат классификатора)
    K — Дата начала действия договора   (из реестра договоров, AL)
    L — Дата окончания договора         (из реестра договоров, AM)
    M — Длящийся (да/нет)               (K и L из разных лет)
    N — Счётчик                         (всегда 1)
    O — Публикация на ЕИС               (правило из шапки примера)
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from datetime import date
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook

from app.config import settings
from app.services.registry_parser import PaymentRow, ContractRow
from app.services.smsp_classifier import (
    ClassifierRule,
    classify_exclusion,
    determine_publication,
    is_continuing,
)
from app.utils.numbers import normalize_text


HEADERS = [
    "Дата",
    "Контрагент",
    "Номер договора",
    "Дата договора",
    "Сумма по договору",
    "Сумма платежа",
    "Вид СМСП",
    "Закупка для СМСП",
    "Способ закупки",
    "Исключение из СМСП",
    "Дата начала действия договора",
    "Дата окончания договора",
    "Длящийся (да/нет)",
    "Счетчик",
    "Публикация на ЕИС (да/нет)",
]


@dataclass(slots=True)
class SumupRow:
    # A..O. Даты сериализуются как ISO-строки при передаче на фронт.
    date: Optional[date]
    contragent: str
    contract_number: str
    contract_date: Optional[date]
    contract_sum: Optional[float]
    payment_sum: float
    smsp_type: str
    smsp_purchase: str
    purchase_method: str
    exclusion_category: str
    action_from: Optional[date]
    action_to: Optional[date]
    is_continuing: str
    counter: int
    publication: str
    # Мета — сколько платежей слилось, есть ли сопоставление с договором
    payments_count: int = 0
    has_contract_match: bool = False

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        for k in ("date", "contract_date", "action_from", "action_to"):
            v = d.get(k)
            if isinstance(v, date):
                d[k] = v.isoformat()
        return d


def build_primary_sumup(
    payments: list[PaymentRow],
    contracts: list[ContractRow],
    rules: list[ClassifierRule],
) -> list[SumupRow]:
    """
    Группирует платежи по договорам и подтягивает реквизиты.

    Ключ группировки (и lookup'а в реестр договоров):
        (normalize(contract_number), contract_date, normalize(contragent))
    При пустом / б/н номере — ключ строится только по (дата, контрагент).
    """
    contract_index = _index_contracts(contracts)

    groups: dict[tuple, list[PaymentRow]] = {}
    for p in payments:
        key = _group_key(p.contract_number, p.contract_date, p.contragent)
        groups.setdefault(key, []).append(p)

    rows: list[SumupRow] = []
    for key, group in groups.items():
        total_rub = sum(p.payment_sum or 0.0 for p in group)
        first_date = min((p.payment_date for p in group if p.payment_date), default=None)
        sample = group[0]

        contract = contract_index.get(key)
        # Fallback: по контрагенту + номеру (без даты), если по точному ключу нет.
        if contract is None:
            contract = _lookup_contract_soft(contracts, sample)

        if contract is not None:
            smsp_type = contract.smsp_type or sample.smsp_type
            smsp_purchase = contract.smsp_purchase or sample.smsp_purchase
            purchase_method = contract.purchase_method
            exclusion_text = contract.exclusion_from_smsp
            action_from = contract.action_from
            action_to = contract.action_to
            contract_sum = contract.contract_sum if contract.contract_sum is not None else sample.contract_sum
        else:
            smsp_type = sample.smsp_type
            smsp_purchase = sample.smsp_purchase
            purchase_method = ""
            exclusion_text = ""
            action_from = None
            action_to = None
            contract_sum = sample.contract_sum

        exclusion_category = classify_exclusion(exclusion_text, rules)
        publication = determine_publication(
            exclusion_category=exclusion_category,
            contract_sum=contract_sum,
            smsp_purchase=smsp_purchase,
            purchase_method=purchase_method,
        )
        continuing = is_continuing(action_from, action_to)

        rows.append(SumupRow(
            date=first_date,
            contragent=sample.contragent,
            contract_number=sample.contract_number,
            contract_date=sample.contract_date,
            contract_sum=contract_sum,
            payment_sum=round(total_rub, 2),
            smsp_type=smsp_type or "",
            smsp_purchase=smsp_purchase or "",
            purchase_method=purchase_method or "",
            exclusion_category=exclusion_category,
            action_from=action_from,
            action_to=action_to,
            is_continuing=continuing,
            counter=1,
            publication=publication,
            payments_count=len(group),
            has_contract_match=contract is not None,
        ))

    # Сортируем как в примере: по дате, затем по контрагенту.
    rows.sort(key=lambda r: (r.date or date.min, r.contragent.lower()))
    return rows


def generate_primary_sumup_xlsx(rows: list[SumupRow], task_id: str) -> str:
    """Записывает первичный свод в xlsx. Возвращает путь к файлу."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Первичный свод"

    # Шапка
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    for i, row in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=row.date)
        ws.cell(row=i, column=2, value=row.contragent)
        ws.cell(row=i, column=3, value=row.contract_number)
        ws.cell(row=i, column=4, value=row.contract_date)
        ws.cell(row=i, column=5, value=row.contract_sum)
        ws.cell(row=i, column=6, value=row.payment_sum)
        ws.cell(row=i, column=7, value=row.smsp_type)
        ws.cell(row=i, column=8, value=row.smsp_purchase)
        ws.cell(row=i, column=9, value=row.purchase_method)
        ws.cell(row=i, column=10, value=row.exclusion_category)
        ws.cell(row=i, column=11, value=row.action_from)
        ws.cell(row=i, column=12, value=row.action_to)
        ws.cell(row=i, column=13, value=row.is_continuing)
        ws.cell(row=i, column=14, value=row.counter)
        ws.cell(row=i, column=15, value=row.publication)

    # Формат дат
    for date_col in (1, 4, 11, 12):
        for r in range(2, len(rows) + 2):
            cell = ws.cell(row=r, column=date_col)
            if cell.value is not None:
                cell.number_format = "DD.MM.YYYY"

    _auto_fit(ws)

    output_dir = settings.GENERATED_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"primary_sumup_{task_id}.xlsx"
    wb.save(str(out_path))
    wb.close()
    return str(out_path)


# ─── Helpers ───────────────────────────────────────────────────────────


def _group_key(
    contract_number: str,
    contract_date: Optional[date],
    contragent: str,
) -> tuple:
    num = normalize_text(contract_number)
    if not num or num == "б н":
        num = None
    return (num, contract_date, normalize_text(contragent))


def _index_contracts(contracts: list[ContractRow]) -> dict[tuple, ContractRow]:
    idx: dict[tuple, ContractRow] = {}
    for c in contracts:
        key = _group_key(c.number, c.conclusion_date, c.contragent)
        idx[key] = c
    return idx


def _lookup_contract_soft(
    contracts: list[ContractRow],
    payment: PaymentRow,
) -> Optional[ContractRow]:
    """Мягкий fallback-поиск: по номеру + контрагенту (без даты)."""
    if not payment.contract_number or normalize_text(payment.contract_number) in ("", "б н"):
        return None
    target_num = normalize_text(payment.contract_number)
    target_con = normalize_text(payment.contragent)
    for c in contracts:
        if normalize_text(c.number) == target_num and normalize_text(c.contragent) == target_con:
            return c
    return None


def _auto_fit(ws) -> None:
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 3, 40)
