"""
Генерация заполненных XLS-шаблонов.

Заполняет 7 столбцов (A-G):
A: Номер документа
B: Дата подписания документа
C: № позиции договора
D: Количество (объем)
E: Единица измерения
F: Цена за единицу с НДС, руб.
G: Сумма по позиции
"""
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from app.config import settings
from app.schemas import MatchResult

YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


async def generate_filled_template(
    matches: list[MatchResult],
    template_path: Optional[str] = None,
    task_id: str = "",
    doc_number: str = "",
    doc_date: str = "",
    merge_duplicates: bool = True,
) -> str:
    """
    Генерирует заполненный XLS-файл на основе результатов сопоставления.

    Args:
        matches: Результаты сопоставления позиций
        template_path: Путь к шаблону (если None — создаёт новый)
        task_id: ID задачи для имени файла
        doc_number: Номер документа (из УПД)
        doc_date: Дата подписания
        merge_duplicates: Объединять ли одинаковые позиции

    Returns:
        Путь к сгенерированному файлу
    """
    if template_path and Path(template_path).exists():
        wb = load_workbook(template_path)
        ws = wb.active
        # Find the first empty row after headers
        start_row = _find_first_empty_row(ws)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Заполненный шаблон"
        # Write headers
        headers = [
            "Номер документа",
            "Дата подписания документа",
            "№ позиции договора",
            "Количество (объем)",
            "Единица измерения",
            "Цена за единицу с НДС, руб.",
            "Сумма по позиции",
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        start_row = 2

    # Optionally merge duplicates (same position number + same price)
    if merge_duplicates:
        matches = _merge_duplicate_matches(matches)

    matches = sorted(
        matches,
        key=lambda item: (
            item.matched_position_number is None,
            item.matched_position_number or 0,
            item.upd_row_index,
        ),
    )

    # Fill rows
    for i, match in enumerate(matches):
        row = start_row + i
        quantity = match.quantity or 1
        price = match.price
        total = match.total or (price * quantity if price else None)

        # If no price from match, calculate from total/quantity
        if not price and total and quantity:
            price = total / quantity

        # Column A: Номер документа
        ws.cell(row=row, column=1, value=doc_number)

        # Column B: Дата подписания документа
        ws.cell(row=row, column=2, value=doc_date)

        # Column C: № позиции договора
        ws.cell(row=row, column=3, value=match.matched_position_number)

        # Column D: Количество (объем)
        ws.cell(row=row, column=4, value=quantity)

        # Column E: Единица измерения
        ws.cell(row=row, column=5, value=match.unit or "шт.")

        # Column F: Цена за единицу с НДС, руб.
        price_cell = ws.cell(row=row, column=6, value=price)

        # Подсветка используется для специальных кейсов, требующих ручной проверки.
        if match.highlight_price or (price is not None and not _is_round_number(price)):
            price_cell.fill = YELLOW_FILL

        # Column G: Сумма по позиции
        ws.cell(row=row, column=7, value=total)

    # Auto-fit column widths
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

    # Save
    output_dir = settings.GENERATED_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"filled_{task_id}.xlsx"
    wb.save(str(output_path))
    wb.close()

    return str(output_path)


def _find_first_empty_row(ws) -> int:
    """Находит первую пустую строку после заголовков."""
    for row in range(2, ws.max_row + 2):
        if all(ws.cell(row=row, column=col).value is None for col in range(1, 8)):
            return row
    return ws.max_row + 1


def _is_round_number(value: float) -> bool:
    """Проверяет, является ли число целым (без копеек)."""
    return abs(value - round(value)) < 0.001


def _merge_duplicate_matches(matches: list[MatchResult]) -> list[MatchResult]:
    """
    Объединяет одинаковые позиции (совпадение номера позиции и цены).
    Объёмы суммируются.

    Дополнительно: если для одной и той же matched_position_number после слияния
    остались строки с разными ценами — все такие строки помечаются needs_review
    + highlight_price с пояснением. Это сигнал для оператора, что часть цен,
    например, попала под коэффициент 1.5, а часть — нет, и нужно убедиться,
    что разделение корректно.
    """
    merged: dict[tuple, MatchResult] = {}
    for m in matches:
        key = (m.matched_position_number, m.price)
        if key in merged:
            existing = merged[key]
            new_qty = (existing.quantity or 0) + (m.quantity or 0)
            new_total = (existing.total or 0) + (m.total or 0)
            merged[key] = MatchResult(
                upd_row_index=existing.upd_row_index,
                upd_item_name=existing.upd_item_name,
                matched_position_number=existing.matched_position_number,
                matched_position_name=existing.matched_position_name,
                confidence=existing.confidence,
                needs_review=existing.needs_review or m.needs_review,
                quantity=new_qty,
                unit=existing.unit,
                price=existing.price,
                total=new_total,
                highlight_price=existing.highlight_price or m.highlight_price,
                highlight_reason=existing.highlight_reason or m.highlight_reason,
            )
        else:
            merged[key] = m

    result = list(merged.values())
    return _flag_position_with_multiple_prices(result)


def _flag_position_with_multiple_prices(matches: list[MatchResult]) -> list[MatchResult]:
    """
    Если matched_position_number встречается с разными ценами — помечаем все такие
    строки needs_review=True и highlight_price=True. matched_position_number=None
    игнорируется (несопоставленные позиции не группируются).
    """
    prices_per_position: dict[int, set[float | None]] = {}
    for m in matches:
        if m.matched_position_number is None:
            continue
        prices_per_position.setdefault(m.matched_position_number, set()).add(m.price)

    flagged_positions = {
        pos for pos, prices in prices_per_position.items() if len(prices) > 1
    }
    if not flagged_positions:
        return matches

    new_reason = "Несколько цен на одну позицию"
    flagged: list[MatchResult] = []
    for m in matches:
        if m.matched_position_number in flagged_positions:
            combined_reason = (
                f"{m.highlight_reason}; {new_reason}" if m.highlight_reason else new_reason
            )
            flagged.append(m.model_copy(update={
                "needs_review": True,
                "highlight_price": True,
                "highlight_reason": combined_reason,
            }))
        else:
            flagged.append(m)
    return flagged
