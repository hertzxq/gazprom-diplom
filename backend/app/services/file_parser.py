"""
Парсинг документов: XLS/XLSX и PDF.

Единый интерфейс: parse_document(path) -> ParsedDocument
"""
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import re
from typing import Any

import openpyxl
import pdfplumber

try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False

HEADER_KEYWORDS = (
    "наименование",
    "услуг",
    "товар",
    "описание",
    "колич",
    "объем",
    "сумм",
    "стоим",
    "цена",
    "ед",
    "номер",
    "пози",
    "счет-фактура",
)

DATE_PATTERN = re.compile(r"\b\d{2}\.\d{2}\.\d{4}\b")
DOC_NUMBER_PATTERNS = (
    re.compile(
        r"(?:счет[\s-]?фактура|сч[её]т[\s-]?фактура)\s*№?\s*([A-Za-zА-Яа-я0-9][A-Za-zА-Яа-я0-9./-]*)",
        re.IGNORECASE,
    ),
    re.compile(
        r"расшифровка к акту[^\n№]{0,80}№\s*([A-Za-zА-Яа-я0-9][A-Za-zА-Яа-я0-9./-]*)",
        re.IGNORECASE,
    ),
    re.compile(
        r"\bакт[^\n№]{0,80}№\s*([A-Za-zА-Яа-я0-9][A-Za-zА-Яа-я0-9./-]*)",
        re.IGNORECASE,
    ),
)


@dataclass(slots=True)
class ParsedDocument:
    rows: list[dict[str, Any]]
    text: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)


def parse_document(file_path: str) -> ParsedDocument:
    """Определяет формат файла и вызывает соответствующий парсер."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".xls":
        if not XLRD_AVAILABLE:
            raise ValueError(
                "Для чтения .xls файлов необходим пакет xlrd. "
                "Установите: pip install xlrd"
            )
        return _parse_xls_xlrd(path)
    if ext == ".xlsx":
        return parse_excel(path)
    if ext == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"Неподдерживаемый формат файла: {ext}")


def _parse_xls_xlrd(path: Path) -> ParsedDocument:
    """Парсит старый .xls формат через xlrd."""
    workbook = xlrd.open_workbook(str(path))
    parsed_rows: list[dict[str, Any]] = []
    text_parts: list[str] = []
    header_detections: list[dict[str, Any]] = []

    for sheet in workbook.sheets():
        if sheet.nrows == 0:
            continue

        raw_rows = []
        for row_idx in range(sheet.nrows):
            row_values = []
            for col_idx in range(sheet.ncols):
                cell = sheet.cell(row_idx, col_idx)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        dt = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                        row_values.append(dt.strftime("%d.%m.%Y"))
                    except Exception:
                        row_values.append(cell.value)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row_values.append(None)
                else:
                    row_values.append(cell.value)
            raw_rows.append(tuple(row_values))

        header_idx, header_score, header_hits = _detect_header_row_with_score(raw_rows)
        header_detections.append({
            "sheet": sheet.name,
            "row": header_idx,
            "score": header_score,
            "keyword_hits": header_hits,
            "confident": header_hits > 0,
        })
        headers = _build_headers(raw_rows[header_idx]) if raw_rows else []

        for line in raw_rows[: min(len(raw_rows), 60)]:
            joined = " ".join(str(cell).strip() for cell in line if cell not in (None, ""))
            if joined:
                text_parts.append(joined)

        for row in raw_rows[header_idx + 1:]:
            if _is_empty_row(row):
                continue
            row_dict: dict[str, Any] = {
                headers[col_idx]: cell
                for col_idx, cell in enumerate(row[: len(headers)])
            }
            row_dict["_row_index"] = len(parsed_rows)
            row_dict["_sheet"] = sheet.name
            parsed_rows.append(row_dict)

    text = "\n".join(text_parts)
    metadata = _extract_document_metadata(text, parsed_rows)
    metadata["header_detection"] = header_detections
    return ParsedDocument(rows=parsed_rows, text=text, metadata=metadata)


def parse_excel(path: Path) -> ParsedDocument:
    """
    Парсит Excel-файл (.xlsx) и извлекает табличные строки вместе с текстовым контекстом.
    """
    workbook = openpyxl.load_workbook(str(path), data_only=True)
    parsed_rows: list[dict[str, Any]] = []
    text_parts: list[str] = []
    header_detections: list[dict[str, Any]] = []

    try:
        for sheet in workbook.worksheets:
            raw_rows = list(sheet.iter_rows(values_only=True))
            if not raw_rows:
                continue

            header_idx, header_score, header_hits = _detect_header_row_with_score(raw_rows)
            header_detections.append({
                "sheet": sheet.title,
                "row": header_idx,
                "score": header_score,
                "keyword_hits": header_hits,
                "confident": header_hits > 0,
            })
            headers = _build_headers(raw_rows[header_idx]) if raw_rows else []

            for line in raw_rows[: min(len(raw_rows), 60)]:
                joined = " ".join(str(cell).strip() for cell in line if cell not in (None, ""))
                if joined:
                    text_parts.append(joined)

            for row in raw_rows[header_idx + 1 :]:
                if _is_empty_row(row):
                    continue
                row_dict: dict[str, Any] = {
                    headers[col_idx]: cell
                    for col_idx, cell in enumerate(row[: len(headers)])
                }
                row_dict["_row_index"] = len(parsed_rows)
                row_dict["_sheet"] = sheet.title
                parsed_rows.append(row_dict)
    finally:
        workbook.close()

    text = "\n".join(text_parts)
    metadata = _extract_document_metadata(text, parsed_rows)
    metadata["header_detection"] = header_detections
    return ParsedDocument(rows=parsed_rows, text=text, metadata=metadata)


def parse_pdf(path: Path) -> ParsedDocument:
    """
    Парсит PDF-файл. Извлекает таблицы и общий текст документа.
    """
    all_rows: list[dict[str, Any]] = []
    text_parts: list[str] = []
    header_detections: list[dict[str, Any]] = []

    with pdfplumber.open(str(path)) as pdf:
        for page_num, page in enumerate(pdf.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)

            tables = page.extract_tables() or []
            for table_idx, table in enumerate(tables):
                if not table or len(table) < 2:
                    continue

                header_idx, header_score, header_hits = _detect_header_row_with_score(table)
                header_detections.append({
                    "page": page_num + 1,
                    "table": table_idx,
                    "row": header_idx,
                    "score": header_score,
                    "keyword_hits": header_hits,
                    "confident": header_hits > 0,
                })
                headers = _build_headers(table[header_idx])

                for row in table[header_idx + 1 :]:
                    if _is_empty_row(row):
                        continue
                    row_dict = {
                        headers[col_idx]: (cell.strip() if isinstance(cell, str) else cell)
                        for col_idx, cell in enumerate(row[: len(headers)])
                    }
                    row_dict["_row_index"] = len(all_rows)
                    row_dict["_page"] = page_num + 1
                    all_rows.append(row_dict)

    text = "\n".join(text_parts)
    metadata = _extract_document_metadata(text, all_rows)
    metadata["header_detection"] = header_detections
    return ParsedDocument(rows=all_rows, text=text, metadata=metadata)


def extract_text_from_pdf(path: Path) -> str:
    """Извлекает весь текст из PDF (для OCR fallback)."""
    text_parts = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
    return "\n".join(text_parts)


def _detect_header_row(rows: list[tuple[Any, ...]] | list[list[Any]]) -> int:
    """
    Возвращает индекс наиболее вероятной строки-заголовка.

    Поведение:
    - Сканируется до 100 строк (раньше было 20) — реальные акты иногда содержат
      длинную преамбулу с реквизитами.
    - Если ни одна строка не содержит ключевых слов из HEADER_KEYWORDS,
      возвращается индекс первой непустой строки (а не 0 «по умолчанию»).
    """
    return _detect_header_row_with_score(rows)[0]


def _detect_header_row_with_score(
    rows: list[tuple[Any, ...]] | list[list[Any]],
) -> tuple[int, int, int]:
    """
    Возвращает (best_index, best_score, keyword_hits).
    Используется напрямую при сборке metadata, чтобы не сканировать дважды.
    """
    best_index = -1
    best_score = -1
    best_keyword_hits = 0
    first_nonempty = -1
    scan_window = min(len(rows), 100)

    for index, row in enumerate(rows[:scan_window]):
        if _is_empty_row(row):
            continue

        if first_nonempty < 0:
            first_nonempty = index

        values = [str(cell).strip().lower() for cell in row if cell not in (None, "")]
        joined = " ".join(values)
        keyword_hits = sum(1 for keyword in HEADER_KEYWORDS if keyword in joined)
        score = keyword_hits * 10 + len(values)

        if score > best_score:
            best_score = score
            best_index = index
            best_keyword_hits = keyword_hits

    if best_keyword_hits == 0:
        # Нет ни одного ключевого слова заголовка — fallback на первую непустую строку.
        if first_nonempty >= 0:
            return first_nonempty, best_score if best_score > 0 else 0, 0
        return 0, 0, 0

    return best_index, best_score, best_keyword_hits


def _build_headers(row: tuple[Any, ...] | list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []

    for col_idx, cell in enumerate(row):
        base_header = str(cell).strip() if cell not in (None, "") else f"col_{col_idx + 1}"
        duplicate_index = seen.get(base_header, 0)
        seen[base_header] = duplicate_index + 1
        if duplicate_index:
            headers.append(f"{base_header}_{duplicate_index + 1}")
        else:
            headers.append(base_header)

    return headers


def _extract_document_metadata(text: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    metadata: dict[str, Any] = {
        "document_number": None,
        "document_date": None,
        "all_dates": [],
    }

    search_text_parts = [text]
    if rows:
        row_preview = []
        for row in rows[:30]:
            joined = " ".join(
                str(value).strip()
                for key, value in row.items()
                if not str(key).startswith("_") and value not in (None, "")
            )
            if joined:
                row_preview.append(joined)
        search_text_parts.append("\n".join(row_preview))

    search_text = "\n".join(part for part in search_text_parts if part).strip()

    for pattern in DOC_NUMBER_PATTERNS:
        match = pattern.search(search_text)
        if match:
            metadata["document_number"] = match.group(1).strip().rstrip(".,;:")
            break

    found_dates: list[str] = []
    for value in DATE_PATTERN.findall(search_text):
        if value not in found_dates:
            found_dates.append(value)

    metadata["all_dates"] = found_dates
    if found_dates:
        metadata["document_date"] = max(
            found_dates,
            key=lambda item: datetime.strptime(item, "%d.%m.%Y"),
        )

    return metadata


def _is_empty_row(row: tuple[Any, ...] | list[Any]) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)
