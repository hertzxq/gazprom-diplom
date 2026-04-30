"""
Второй свод СМСП (Задача 2, шаг 2).

Строит xlsx с 3 вкладками по данным первичного свода. Формулы — из шапки
`examples/Task2/Второй свод_обобщение.xlsx`, пересказанной на языке колонок
первичного свода (A–O):

    F — Сумма платежа
    G — Вид СМСП        (малое/микро/самозанятый/среднее/нет)
    H — Закупка для СМСП (да/нет)
    J — Исключение из СМСП (авиа/аренда/страх/образов/почта/нет)
    O — Публикация на ЕИС (да/нет)

27 строк показателей на каждой вкладке:

    2   Итого всего                                          (все строки)
    4–8 Исключения по категориям                             (J == cat)
    9   Аренда СМСП   (J == аренда) И (G != «нет»)
    10  СМСП образование (J == образов) И (G != «нет»)
    12  Итого за минусом исключений = строка2 − Σ(строки 4–8)
    14–16 «Только для СМСП»:
        14  H == «да»
        15  H == «да» И G == «среднее»
        16  H == «да» И G ∈ {«малое», «микро», «самозанятый»}
    18–20 «С СМСП для всех»:
        18  H == «нет» И G != «нет»
        19  H == «нет» И G == «среднее»
        20  H == «нет» И G ∈ {«малое», «микро», «самозанятый»}
    22  строка14 + строка18
    24–26 как 18–20, но + J == «нет»
    27  строка14 + строка24

Вкладки:
    Итого_платежи_договоры            — все строки
    Итого_платежи_договоры_публ       — только O == «да»
    Итого_платежи_договоры_не_публ    — только O == «нет»
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import settings
from app.services.primary_sumup import SumupRow
from app.utils.numbers import normalize_text


SMSP_SMALL = {"малое", "микро", "самозанятый"}
SMSP_MEDIUM = {"среднее"}

# Список категорий исключений, которые отображаются в блоке «Исключения».
# Порядок важен — он же отражается в строках 4–8 итогового свода.
EXCLUSION_CATEGORIES_DISPLAY: list[tuple[str, str]] = [
    ("Исключения (авиа)",     "авиа"),
    ("Исключение (страх)",    "страх"),
    ("Исключение (образов)",  "образов"),
    ("Исключение (аренда)",   "аренда"),
    ("Исключение (почта)",    "почта"),
]

SHEETS: list[tuple[str, Callable[[SumupRow], bool]]] = [
    ("Итого_платежи_договоры",          lambda r: True),
    ("Итого_платежи_договоры_публ",     lambda r: _norm(r.publication) == "да"),
    ("Итого_платежи_договоры_не_публ",  lambda r: _norm(r.publication) == "нет"),
]


@dataclass(slots=True)
class MetricRow:
    label: str
    total_sum: float
    count: Optional[int]  # None → в ячейку не пишем (служебные «итоговые» строки)


def generate_final_summary_xlsx(rows: list[SumupRow], task_id: str) -> str:
    wb = Workbook()
    # Удалим дефолтный лист — создадим по списку SHEETS.
    wb.remove(wb.active)

    for sheet_name, predicate in SHEETS:
        subset = [r for r in rows if predicate(r)]
        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, subset)

    output_dir = settings.GENERATED_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"smsp_final_summary_{task_id}.xlsx"
    wb.save(str(out_path))
    wb.close()
    return str(out_path)


def compute_final_summary(rows: list[SumupRow]) -> dict[str, list[dict]]:
    """
    Возвращает результаты 3 вкладок в виде списков метрик — для preview
    на фронте (шаг 3) без необходимости скачивать xlsx.
    """
    result: dict[str, list[dict]] = {}
    for sheet_name, predicate in SHEETS:
        subset = [r for r in rows if predicate(r)]
        metrics = _build_metrics(subset)
        result[sheet_name] = [
            {"label": m.label, "total_sum": m.total_sum, "count": m.count}
            for m in metrics if m is not None
        ]
    return result


# ─── Internals ─────────────────────────────────────────────────────────


def _write_sheet(ws, subset: list[SumupRow]) -> None:
    # Шапка — как в `examples/Task2/Второй свод_обобщение.xlsx`.
    header = ["Наименование показателя", "Сумма", "Количество договоров", "Комментарий"]
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E8EEF7", end_color="E8EEF7", fill_type="solid")
    for col, value in enumerate(header, start=1):
        c = ws.cell(row=1, column=col, value=value)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")

    metrics = _build_metrics(subset)
    # Раскладка по строкам 2..27 как в примере. _build_metrics возвращает
    # уже в нужном порядке (с None-заглушками для разделителей).
    for i, m in enumerate(metrics, start=2):
        if m is None:
            continue
        ws.cell(row=i, column=1, value=m.label)
        ws.cell(row=i, column=2, value=round(m.total_sum, 2))
        if m.count is not None:
            ws.cell(row=i, column=3, value=m.count)

    _auto_fit(ws)


def _build_metrics(rows: list[SumupRow]) -> list[Optional[MetricRow]]:
    """
    Строит упорядоченный список метрик (с None-разделителями на пустых строках).

    Индексы соответствуют строкам xlsx (index 0 → строка 2, index 1 → строка 3, …).
    """
    metrics: list[Optional[MetricRow]] = []

    # Строка 2 — «Итого всего»
    metrics.append(_total(rows, "Итого всего"))
    metrics.append(None)  # 3 — пусто

    # Строки 4–8 — исключения (авиа/страх/образов/аренда/почта)
    for label, category in EXCLUSION_CATEGORIES_DISPLAY:
        metrics.append(_filter_metric(
            rows, label,
            lambda r, c=category: _norm(r.exclusion_category) == c,
        ))

    # Строка 9 — аренда+СМСП
    metrics.append(_filter_metric(
        rows, "Аренда СМСП",
        lambda r: _norm(r.exclusion_category) == "аренда" and _norm(r.smsp_type) != "нет" and _norm(r.smsp_type) != "",
    ))
    # Строка 10 — СМСП образование
    metrics.append(_filter_metric(
        rows, "СМСП образование",
        lambda r: _norm(r.exclusion_category) == "образов" and _norm(r.smsp_type) != "нет" and _norm(r.smsp_type) != "",
    ))
    metrics.append(None)  # 11

    # Строка 12 — Итого за минусом исключений (2 − Σ 4–8)
    total_all = metrics[0].total_sum if metrics[0] else 0.0
    total_excl_sum = sum(
        metrics[i].total_sum for i in range(2, 7)  # индексы 2..6 → строки 4..8
        if metrics[i] is not None
    )
    metrics.append(MetricRow("Итого за минусом исключений", total_all - total_excl_sum, count=None))
    metrics.append(None)  # 13

    # Строки 14–16 — «Только для СМСП» (H = «да»)
    metrics.append(_filter_metric(
        rows, "Только для СМСП",
        lambda r: _norm(r.smsp_purchase) == "да",
    ))
    metrics.append(_filter_metric(
        rows, "Только для СМСП среднее",
        lambda r: _norm(r.smsp_purchase) == "да" and _norm(r.smsp_type) in SMSP_MEDIUM,
    ))
    metrics.append(_filter_metric(
        rows, "Только для СМСП малое и микро",
        lambda r: _norm(r.smsp_purchase) == "да" and _norm(r.smsp_type) in SMSP_SMALL,
    ))
    metrics.append(None)  # 17

    # Строки 18–20 — «С СМСП для всех» (H = «нет» И G != «нет»/пусто)
    metrics.append(_filter_metric(
        rows, "С СМСП для всех",
        lambda r: _norm(r.smsp_purchase) == "нет"
                  and _norm(r.smsp_type) not in ("", "нет"),
    ))
    metrics.append(_filter_metric(
        rows, "с СМСП среднее",
        lambda r: _norm(r.smsp_purchase) == "нет" and _norm(r.smsp_type) in SMSP_MEDIUM,
    ))
    metrics.append(_filter_metric(
        rows, "с СМСП малое и микро",
        lambda r: _norm(r.smsp_purchase) == "нет" and _norm(r.smsp_type) in SMSP_SMALL,
    ))
    metrics.append(None)  # 21

    # Строка 22 — Итого с СМСП = 14 + 18
    only_smsp = metrics[12].total_sum if metrics[12] else 0.0
    with_smsp = metrics[16].total_sum if metrics[16] else 0.0
    metrics.append(MetricRow("Итого с СМСП", only_smsp + with_smsp, count=None))
    metrics.append(None)  # 23

    # Строки 24–26 — 18–20 с дополнительным фильтром J == «нет»
    metrics.append(_filter_metric(
        rows, "с СМСП для всех (минус исключения)",
        lambda r: _norm(r.smsp_purchase) == "нет"
                  and _norm(r.smsp_type) not in ("", "нет")
                  and _norm(r.exclusion_category) == "нет",
    ))
    metrics.append(_filter_metric(
        rows, "с СМСП для всех среднее (минус исключения)",
        lambda r: _norm(r.smsp_purchase) == "нет"
                  and _norm(r.smsp_type) in SMSP_MEDIUM
                  and _norm(r.exclusion_category) == "нет",
    ))
    metrics.append(_filter_metric(
        rows, "с СМСП малое и микро (минус исключения)",
        lambda r: _norm(r.smsp_purchase) == "нет"
                  and _norm(r.smsp_type) in SMSP_SMALL
                  and _norm(r.exclusion_category) == "нет",
    ))

    # Строка 27 — Итого с СМСП за минусом исключений = 14 + 24
    only_smsp_minus = metrics[12].total_sum if metrics[12] else 0.0
    with_smsp_minus = metrics[22].total_sum if metrics[22] else 0.0
    metrics.append(MetricRow(
        "Итого с СМСП (за минусом исключений)",
        only_smsp_minus + with_smsp_minus,
        count=None,
    ))

    return metrics


def _total(rows: list[SumupRow], label: str) -> MetricRow:
    return MetricRow(
        label=label,
        total_sum=sum(r.payment_sum for r in rows),
        count=len(rows),
    )


def _filter_metric(
    rows: list[SumupRow],
    label: str,
    predicate: Callable[[SumupRow], bool],
) -> MetricRow:
    filtered = [r for r in rows if predicate(r)]
    return MetricRow(
        label=label,
        total_sum=sum(r.payment_sum for r in filtered),
        count=len(filtered),
    )


def _norm(value: Optional[str]) -> str:
    return normalize_text(value or "")


def _auto_fit(ws) -> None:
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 3, 50)
