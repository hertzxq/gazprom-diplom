"""API-эндпоинты для поиска производителей.

Поиск выполняется LLM на CPU и занимает минуты, поэтому он вынесен из
HTTP-запроса в фоновую asyncio-задачу: POST сразу возвращает task_id,
фронт опрашивает GET /tasks/{id} и может отменить поиск POST /tasks/{id}/cancel.
Так обновление страницы или уход с неё не теряют результат, а отмена
реально обрывает соединение с Ollama (генерация прекращается).
"""
import asyncio
import os
import uuid
import tempfile
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.background import BackgroundTask

from app.database import get_db, async_session
from app.models.user import User
from app.models.document import ProcessingTask, DocumentStatus
from app.schemas import (
    ManufacturerSearchRequest,
    ManufacturerInfoRequest,
    ManufacturerResult,
    ManufacturerDocResult,
    ManufacturerExportRequest,
    ManufacturerTaskStartResponse,
    ManufacturerTaskStatusResponse,
)
from app.services.manufacturer_search import (
    search_manufacturers_by_specs,
    search_manufacturer_info,
)
from app.utils.auth import get_current_user

router = APIRouter(prefix="/api/manufacturers", tags=["Поиск производителей"])

# Живые поиски текущего процесса: str(task_id) → asyncio.Task.
# Бэкенд работает одним процессом, поэтому словаря в памяти достаточно.
_RUNNING: dict[str, asyncio.Task] = {}

CANCELLED_MESSAGE = "Поиск отменён пользователем"


async def _finalize_task(task_id: uuid.UUID, status: DocumentStatus, matching_data: dict) -> None:
    """
    Записывает итог фоновой задачи отдельной сессией — сессия исходного
    HTTP-запроса к этому моменту уже закрыта. PROCESSING-guard защищает
    от гонки «отмена против позднего результата».
    """
    async with async_session() as session:
        task = await session.get(ProcessingTask, task_id)
        if task is not None and task.status == DocumentStatus.PROCESSING:
            task.status = status
            task.matching_data = matching_data
            await session.commit()


def _start_background_search(task_id: uuid.UUID, coro) -> None:
    """Запускает поиск фоном и регистрирует его для статуса/отмены."""
    bg = asyncio.create_task(coro)
    key = str(task_id)
    _RUNNING[key] = bg
    bg.add_done_callback(lambda _t, _k=key: _RUNNING.pop(_k, None))


async def _run_specs_search(task_id: uuid.UUID, request: ManufacturerSearchRequest) -> None:
    try:
        raw_results = await search_manufacturers_by_specs(
            product_name=request.product_name,
            characteristics=[c.model_dump() for c in request.characteristics],
            sources=request.sources,
        )
        results = [
            ManufacturerResult(
                name=r.get("name", "Неизвестно"),
                country=r.get("country"),
                website=r.get("website"),
                contacts=r.get("contacts"),
                certificates=r.get("certificates"),
                products=r.get("products"),
                source=r.get("source"),
            ).model_dump()
            for r in raw_results
        ]
        await _finalize_task(task_id, DocumentStatus.PROCESSED, {
            "results_count": len(results),
            "results": results,
        })
    except asyncio.CancelledError:
        # shield: запись об отмене должна дойти до БД, даже когда задачу отменяют
        await asyncio.shield(_finalize_task(task_id, DocumentStatus.ERROR, {
            "cancelled": True,
            "error": CANCELLED_MESSAGE,
        }))
        raise
    except Exception as e:
        await _finalize_task(task_id, DocumentStatus.ERROR, {"error": str(e)})


async def _run_name_search(task_id: uuid.UUID, request: ManufacturerInfoRequest) -> None:
    try:
        raw = await search_manufacturer_info(
            product_name=request.product_name,
            sources=request.sources,
        )
        manufacturers = [
            ManufacturerResult(
                name=m.get("name", "Неизвестно"),
                country=m.get("country"),
                website=m.get("website"),
                contacts=m.get("contacts"),
                is_primary=m.get("is_primary", False),
            ).model_dump()
            for m in raw.get("manufacturers", [])
        ]
        documentation = [
            ManufacturerDocResult(
                title=d.get("title", "Без названия"),
                doc_type=d.get("doc_type"),
                source_url=d.get("source_url"),
                description=d.get("description"),
            ).model_dump()
            for d in raw.get("documentation", [])
        ]
        await _finalize_task(task_id, DocumentStatus.PROCESSED, {
            "manufacturers_count": len(manufacturers),
            "documentation_count": len(documentation),
            "manufacturers": manufacturers,
            "documentation": documentation,
            "summary": raw.get("summary", ""),
        })
    except asyncio.CancelledError:
        await asyncio.shield(_finalize_task(task_id, DocumentStatus.ERROR, {
            "cancelled": True,
            "error": CANCELLED_MESSAGE,
        }))
        raise
    except Exception as e:
        await _finalize_task(task_id, DocumentStatus.ERROR, {"error": str(e)})


async def _create_search_task(db: AsyncSession, user: User, task_type: str, parameters: dict) -> ProcessingTask:
    task = ProcessingTask(
        task_type=task_type,
        parameters=parameters,
        created_by=user.id,
        status=DocumentStatus.PROCESSING,
    )
    db.add(task)
    # Коммитим сразу: строка должна быть видна опросу /tasks/{id}
    # и фоновой задаче (другие сессии) до завершения этого запроса.
    await db.commit()
    await db.refresh(task)
    return task


@router.post("/search-by-specs", response_model=ManufacturerTaskStartResponse)
async def search_by_specs(
    request: ManufacturerSearchRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Запускает фоновый поиск производителей по характеристикам товара."""
    task = await _create_search_task(
        db, current_user, "manufacturer_search_specs",
        {
            "product_name": request.product_name,
            "characteristics": [c.model_dump() for c in request.characteristics],
            "sources": request.sources,
        },
    )
    _start_background_search(task.id, _run_specs_search(task.id, request))
    return ManufacturerTaskStartResponse(
        task_id=task.id,
        message="Поиск запущен — на CPU это занимает 1–3 минуты",
    )


@router.post("/search-by-name", response_model=ManufacturerTaskStartResponse)
async def search_by_name(
    request: ManufacturerInfoRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Запускает фоновый поиск производителя и документации по наименованию."""
    task = await _create_search_task(
        db, current_user, "manufacturer_search_name",
        {
            "product_name": request.product_name,
            "sources": request.sources,
        },
    )
    _start_background_search(task.id, _run_name_search(task.id, request))
    return ManufacturerTaskStartResponse(
        task_id=task.id,
        message="Поиск запущен — на CPU это занимает 1–3 минуты",
    )


@router.get("/tasks/{task_id}", response_model=ManufacturerTaskStatusResponse)
async def get_search_task(
    task_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Статус и результат фонового поиска (для опроса с фронта)."""
    task = await db.get(ProcessingTask, task_id)
    if task is None or not (task.task_type or "").startswith("manufacturer_search"):
        raise HTTPException(status_code=404, detail="Задача поиска не найдена")

    # PROCESSING в БД при отсутствии живой задачи в процессе = бэкенд
    # перезапускался. Перечитываем строку (защита от гонки с финализацией)
    # и помечаем задачу потерянной.
    if task.status == DocumentStatus.PROCESSING and str(task_id) not in _RUNNING:
        await db.refresh(task)
        if task.status == DocumentStatus.PROCESSING:
            task.status = DocumentStatus.ERROR
            task.matching_data = {"error": "Поиск прерван перезапуском сервера — запустите заново"}
            await db.commit()

    data = task.matching_data or {}
    return ManufacturerTaskStatusResponse(
        task_id=task.id,
        task_type=task.task_type,
        status=task.status.value,
        cancelled=bool(data.get("cancelled")),
        error=data.get("error"),
        results=data.get("results"),
        manufacturers=data.get("manufacturers"),
        documentation=data.get("documentation"),
        summary=data.get("summary"),
    )


@router.post("/tasks/{task_id}/cancel", response_model=ManufacturerTaskStartResponse)
async def cancel_search_task(
    task_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Отменяет выполняющийся поиск. Запись в БД делает сам отменяемый таск."""
    task = await db.get(ProcessingTask, task_id)
    if task is None or not (task.task_type or "").startswith("manufacturer_search"):
        raise HTTPException(status_code=404, detail="Задача поиска не найдена")

    bg = _RUNNING.get(str(task_id))
    if bg is None or bg.done():
        raise HTTPException(status_code=409, detail="Поиск уже завершён")

    bg.cancel()
    return ManufacturerTaskStartResponse(task_id=task_id, message=CANCELLED_MESSAGE)


@router.post("/export")
async def export_results(
    request: ManufacturerExportRequest,
    current_user: User = Depends(get_current_user),
):
    """Экспорт результатов поиска производителей в XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    if not request.results:
        raise HTTPException(status_code=400, detail="Нет данных для экспорта")

    wb = Workbook()
    ws = wb.active
    ws.title = "Производители"

    # Header style
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="003E92", end_color="003E92", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["№", "Производитель", "Страна", "Продукция", "Сертификаты", "Контакты", "Сайт", "Источник"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    data_alignment = Alignment(vertical="top", wrap_text=True)
    for row_idx, r in enumerate(request.results, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1).alignment = data_alignment
        ws.cell(row=row_idx, column=2, value=r.name).alignment = data_alignment
        ws.cell(row=row_idx, column=3, value=r.country or "").alignment = data_alignment
        ws.cell(row=row_idx, column=4, value=r.products or "").alignment = data_alignment
        ws.cell(row=row_idx, column=5, value=r.certificates or "").alignment = data_alignment
        ws.cell(row=row_idx, column=6, value=r.contacts or "").alignment = data_alignment
        ws.cell(row=row_idx, column=7, value=r.website or "").alignment = data_alignment
        ws.cell(row=row_idx, column=8, value=r.source or "").alignment = data_alignment

        for col_idx in range(1, 9):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    # Auto-width
    for col_idx in range(1, 9):
        max_length = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, len(request.results) + 2)
        )
        ws.column_dimensions[chr(64 + col_idx)].width = min(max(max_length + 2, 10), 40)

    # Save to temp file
    product_slug = (request.product_name or "results").replace(" ", "_")[:30]
    tmp = tempfile.NamedTemporaryFile(
        delete=False, suffix=".xlsx", prefix=f"manufacturers_{product_slug}_"
    )
    wb.save(tmp.name)
    tmp.close()

    return FileResponse(
        path=tmp.name,
        filename=f"Производители_{product_slug}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        # Файл одноразовый: удаляем после отдачи, чтобы не копить мусор в temp.
        background=BackgroundTask(os.unlink, tmp.name),
    )

