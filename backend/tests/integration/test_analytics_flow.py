"""
Integration: GET /api/analytics/summary возвращает все ожидаемые ключи,
а Dashboard может ими пользоваться без хардкода.
"""
import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

from app.models.document import Document, DocumentStatus, FileType


def _xlsx_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


@pytest.mark.asyncio
class TestAnalyticsSummary:
    async def test_returns_all_expected_keys(
        self, async_client, regular_user, user_token, auth_headers,
    ):
        response = await async_client.get(
            "/api/analytics/summary", headers=auth_headers(user_token),
        )
        assert response.status_code == 200
        data = response.json()
        # Все обязательные ключи присутствуют
        for key in [
            "total", "processed", "pending", "errors",
            "success_rate", "error_rate",
            "by_status", "by_type",
            "monthly", "yearly", "volume_trend", "top_items", "region_status",
        ]:
            assert key in data, f"Missing key: {key}"

    async def test_empty_state(self, async_client, regular_user, user_token, auth_headers):
        response = await async_client.get(
            "/api/analytics/summary", headers=auth_headers(user_token),
        )
        data = response.json()
        assert data["total"] == 0
        assert data["success_rate"] == 0
        assert data["monthly"] == []
        assert data["volume_trend"] == []
        assert data["top_items"] == []

    async def test_counts_uploaded_documents(
        self, async_client, regular_user, user_token, auth_headers, db_session,
    ):
        for i in range(3):
            db_session.add(Document(
                filename=f"f{i}.xlsx",
                original_filename=f"f{i}.xlsx",
                file_type=FileType.POSITIONS,
                file_path=f"/tmp/f{i}.xlsx",
                status=DocumentStatus.UPLOADED,
                uploaded_by=regular_user.id,
            ))
        db_session.add(Document(
            filename="done.xlsx",
            original_filename="done.xlsx",
            file_type=FileType.UPD,
            file_path="/tmp/done.xlsx",
            status=DocumentStatus.PROCESSED,
            uploaded_by=regular_user.id,
        ))
        await db_session.commit()

        response = await async_client.get(
            "/api/analytics/summary", headers=auth_headers(user_token),
        )
        data = response.json()
        assert data["total"] == 4
        assert data["processed"] == 1
        assert data["pending"] == 3
        assert data["success_rate"] == 25.0  # 1/4
        # by_type содержит positions и upd
        assert data["by_type"]["positions"] == 3
        assert data["by_type"]["upd"] == 1
