"""
Unit-тесты парсеров реестров (Задача 2).

Генерируем временные xlsx с известной структурой и сверяем, что парсер
достаёт данные по буквенным индексам колонок, правильно обрабатывает:
    * множественные платежи в одной ячейке (через `;`);
    * неразрывные пробелы в суммах (`\\xa0`);
    * запятую как десятичный разделитель;
    * валютные платежи через колонку AJ (сумма в рублях).
"""
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.services.registry_parser import (
    parse_payment_registry,
    parse_contract_registry,
    _col_letter_to_index,
    _parse_payment_amount,
    _to_date,
)


# ─── Хелперы для генерации тестовых xlsx ───────────────────────────────


def _make_payment_xlsx(path: Path, payment_rows: list[dict]) -> Path:
    """
    Создаёт тестовый реестр платежей. Каждая строка — dict с колонками
    (ключи — буквенные индексы). Пустые колонки заполняются None.
    Шапка — в 3-й строке; данные — с 4-й.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"

    # 1-я строка — заголовок секции (как в реальном файле)
    ws.cell(row=1, column=1, value="Бух. документы")

    # 3-я строка — шапка с названиями колонок
    headers_by_letter = {
        "A": "Вложения", "B": "Важный", "C": "Сейчас у", "D": "Вид",
        "E": "Номер", "F": "Дата", "G": "Контрагент", "H": "Номер договора",
        "I": "Дата договора", "J": "Сумма по договору", "K": "Сумма платежа",
        "M": "Вид СМСП", "N": "Закупка для СМСП", "O": "Валюта по договору",
        "AJ": "Сумма в рублях для валютных платежей", "AK": "Валюта платежа",
    }
    for letter, name in headers_by_letter.items():
        ws.cell(row=3, column=_col_letter_to_index(letter) + 1, value=name)

    # Данные — с 4-й строки
    for i, row_data in enumerate(payment_rows, start=4):
        for letter, value in row_data.items():
            ws.cell(row=i, column=_col_letter_to_index(letter) + 1, value=value)

    wb.save(str(path))
    wb.close()
    return path


def _make_contract_xlsx(path: Path, contract_rows: list[dict]) -> Path:
    """Создаёт тестовый реестр договоров аналогично."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"

    ws.cell(row=1, column=1, value="Договоры")
    headers = {
        "A": "Вид", "B": "Номер", "C": "Дата заключения",
        "E": "Сумма с НДС", "F": "Вид закупки", "G": "Вид СМСП",
        "H": "Закупка для СМСП", "N": "Исключение из СМСП",
        "AL": "Исполнение с", "AM": "Исполнение по", "BM": "Контрагенты",
    }
    for letter, name in headers.items():
        ws.cell(row=3, column=_col_letter_to_index(letter) + 1, value=name)

    for i, row_data in enumerate(contract_rows, start=4):
        for letter, value in row_data.items():
            ws.cell(row=i, column=_col_letter_to_index(letter) + 1, value=value)

    wb.save(str(path))
    wb.close()
    return path


# ─── Колонки и индексы ─────────────────────────────────────────────────


class TestColLetterToIndex:
    def test_single_letter(self):
        assert _col_letter_to_index("A") == 0
        assert _col_letter_to_index("Z") == 25

    def test_double_letter(self):
        assert _col_letter_to_index("AA") == 26
        assert _col_letter_to_index("AJ") == 35
        assert _col_letter_to_index("AK") == 36
        assert _col_letter_to_index("BM") == 64

    def test_lowercase(self):
        assert _col_letter_to_index("aj") == 35


# ─── Парсинг сумм ──────────────────────────────────────────────────────


class TestParsePaymentAmount:
    def test_plain_number(self):
        total, multi = _parse_payment_amount(1000.50)
        assert total == pytest.approx(1000.50)
        assert multi is False

    def test_string_with_nbsp(self):
        total, multi = _parse_payment_amount("8\xa0100")
        assert total == pytest.approx(8100.0)
        assert multi is False

    def test_multiple_payments_split_semicolon(self):
        total, multi = _parse_payment_amount("51\xa0200; 184\xa0000")
        assert total == pytest.approx(235200.0)
        assert multi is True

    def test_russian_decimal(self):
        total, _ = _parse_payment_amount("322\xa0069,36")
        assert total == pytest.approx(322069.36)

    def test_empty(self):
        total, multi = _parse_payment_amount(None)
        assert total is None
        assert multi is False

        total, multi = _parse_payment_amount("")
        assert total is None

    def test_three_semicolons(self):
        total, multi = _parse_payment_amount("1000; 2000; 3000")
        assert total == pytest.approx(6000.0)
        assert multi is True


# ─── Парсинг дат ───────────────────────────────────────────────────────


class TestToDate:
    def test_datetime(self):
        from datetime import datetime
        assert _to_date(datetime(2026, 4, 18)) == date(2026, 4, 18)

    def test_date(self):
        assert _to_date(date(2026, 4, 18)) == date(2026, 4, 18)

    def test_none_and_empty(self):
        assert _to_date(None) is None
        assert _to_date("") is None

    def test_string_iso(self):
        assert _to_date("2026-04-18") == date(2026, 4, 18)

    def test_string_ru(self):
        assert _to_date("18.04.2026") == date(2026, 4, 18)

    def test_excel_serial(self):
        # 44631 = 2022-03-11 в Excel
        d = _to_date(44631.0)
        assert d == date(2022, 3, 11)

    def test_garbage(self):
        assert _to_date("not a date") is None


# ─── Парсинг реестра платежей (через сгенерированный xlsx) ─────────────


class TestParsePaymentRegistry:
    def test_rub_payment(self, tmp_path):
        path = _make_payment_xlsx(tmp_path / "pay.xlsx", [
            {
                "D": "Платеж",
                "F": date(2026, 1, 15),
                "G": "ООО Ромашка",
                "H": "DOG-001",
                "I": date(2026, 1, 10),
                "J": 100_000.0,
                "K": "50\xa0000",
                "M": "Нет",
                "N": "Нет",
                "O": "RUB",
            },
        ])
        rows = parse_payment_registry(str(path))
        assert len(rows) == 1
        r = rows[0]
        assert r.contragent == "ООО Ромашка"
        assert r.contract_number == "DOG-001"
        assert r.contract_date == date(2026, 1, 10)
        assert r.payment_sum == pytest.approx(50_000.0)
        assert r.currency == "RUB"
        assert r.has_multi_payments is False

    def test_currency_payment_uses_aj(self, tmp_path):
        """Для USD-платежа берётся колонка AJ «Сумма в рублях», не K."""
        path = _make_payment_xlsx(tmp_path / "pay.xlsx", [
            {
                "D": "Платеж", "F": date(2022, 3, 11), "G": "Bloomberg",
                "H": "3107460", "I": date(2020, 8, 26), "J": 64800.0,
                "K": "8\xa0100",     # это USD
                "O": "USD",
                "AJ": "848889.72",   # это уже в рублях
            },
        ])
        rows = parse_payment_registry(str(path))
        assert len(rows) == 1
        assert rows[0].payment_sum == pytest.approx(848889.72)
        assert rows[0].currency == "USD"

    def test_multi_payments_in_one_cell(self, tmp_path):
        """Платёж типа `194970; 36702` суммируется в 231672."""
        path = _make_payment_xlsx(tmp_path / "pay.xlsx", [
            {
                "D": "Платеж", "F": date(2025, 4, 22), "G": "Marusia",
                "H": "2025-GTC", "I": date(2025, 4, 22),
                "K": "194\xa0970; 36\xa0702",
                "O": "RUB",
            },
        ])
        rows = parse_payment_registry(str(path))
        assert rows[0].payment_sum == pytest.approx(231672.0)
        assert rows[0].has_multi_payments is True

    def test_skip_non_payment_rows(self, tmp_path):
        """Строки без `Вид='Платеж'` пропускаются (заголовки, пустые)."""
        path = _make_payment_xlsx(tmp_path / "pay.xlsx", [
            {"D": "Платеж", "G": "A", "K": 100, "O": "RUB"},
            {"D": "", "G": "B", "K": 200},              # пустой вид — пропуск
            {"D": "Договор", "G": "C", "K": 300},       # не платёж — пропуск
            {"D": "Платеж", "G": "D", "K": 400, "O": "RUB"},
        ])
        rows = parse_payment_registry(str(path))
        assert len(rows) == 2
        assert [r.contragent for r in rows] == ["A", "D"]


# ─── Парсинг реестра договоров ─────────────────────────────────────────


class TestParseContractRegistry:
    def test_basic(self, tmp_path):
        path = _make_contract_xlsx(tmp_path / "con.xlsx", [
            {
                "A": "Договор", "B": "DOG-001", "C": date(2026, 1, 10),
                "E": 100_000.0, "F": "ЕП", "G": "Нет", "H": "Нет",
                "N": "Не является исключением",
                "AL": date(2026, 1, 15), "AM": date(2026, 12, 31),
                "BM": "ООО Ромашка",
            },
        ])
        rows = parse_contract_registry(str(path))
        assert len(rows) == 1
        r = rows[0]
        assert r.number == "DOG-001"
        assert r.conclusion_date == date(2026, 1, 10)
        assert r.contract_sum == pytest.approx(100_000.0)
        assert r.purchase_method == "ЕП"
        assert r.exclusion_from_smsp == "Не является исключением"
        assert r.action_from == date(2026, 1, 15)
        assert r.action_to == date(2026, 12, 31)
        assert r.contragent == "ООО Ромашка"

    def test_skip_addendum(self, tmp_path):
        """«Дополнительное соглашение» не попадает в первичный свод."""
        path = _make_contract_xlsx(tmp_path / "con.xlsx", [
            {"A": "Договор", "B": "001", "BM": "A"},
            {"A": "Дополнительное соглашение", "B": "001", "BM": "B"},
            {"A": "Договор", "B": "002", "BM": "C"},
        ])
        rows = parse_contract_registry(str(path))
        assert [r.number for r in rows] == ["001", "002"]


# ─── Интеграционные тесты на реальных файлах из examples/Task2/ ────────
# Запускаются только если файлы существуют (в CI могут отсутствовать).


REAL_PAYMENTS = Path(__file__).resolve().parents[2] / "examples" / "Task2" / "Реестр_платежей.xls"
REAL_CONTRACTS = Path(__file__).resolve().parents[2] / "examples" / "Task2" / "Реестр договоров.xls"


@pytest.mark.skipif(not REAL_PAYMENTS.exists(), reason="Реальный реестр платежей недоступен")
class TestRealExamples:
    def test_payment_registry_row_count(self):
        rows = parse_payment_registry(str(REAL_PAYMENTS))
        # В реальном файле 4271 платёж (4278 строк − 3 заголовка − 4 не-платежа).
        # Проверяем, что их много и у них разные валюты.
        assert len(rows) > 4000
        currencies = {r.currency for r in rows}
        assert {"RUB", "USD"}.issubset(currencies)

    def test_contract_registry_has_rental(self):
        rows = parse_contract_registry(str(REAL_CONTRACTS))
        rental = [r for r in rows if "2026-Вб" in r.number]
        assert len(rental) == 1
        assert "аренд" in rental[0].exclusion_from_smsp.lower()
