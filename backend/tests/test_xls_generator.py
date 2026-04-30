"""
Unit-тесты для xls_generator: проверяем генерацию XLS, подсветку, объединение дубликатов.
"""
from pathlib import Path

import openpyxl
import pytest

from app.config import settings
from app.schemas import MatchResult
from app.services.xls_generator import (
    _find_first_empty_row,
    _is_round_number,
    _merge_duplicate_matches,
    generate_filled_template,
)


def _mr(**kwargs) -> MatchResult:
    """Хелпер — MatchResult с дефолтами."""
    defaults = {
        "upd_row_index": 0,
        "upd_item_name": "Услуга",
        "matched_position_number": 1,
        "matched_position_name": "Услуга из прайса",
        "confidence": 90.0,
        "needs_review": False,
        "quantity": 1.0,
        "unit": "шт.",
        "price": 1000.0,
        "total": 1000.0,
        "highlight_price": False,
        "highlight_reason": None,
    }
    defaults.update(kwargs)
    return MatchResult(**defaults)


class TestIsRoundNumber:
    def test_round(self):
        assert _is_round_number(1000.0) is True

    def test_not_round(self):
        assert _is_round_number(1000.5) is False


class TestMergeDuplicates:
    def test_merges_same_number_and_price(self):
        matches = [
            _mr(matched_position_number=1, price=500.0, quantity=2.0, total=1000.0),
            _mr(matched_position_number=1, price=500.0, quantity=3.0, total=1500.0),
            _mr(matched_position_number=2, price=700.0, quantity=1.0, total=700.0),
        ]
        merged = _merge_duplicate_matches(matches)
        assert len(merged) == 2
        first = next(m for m in merged if m.matched_position_number == 1)
        assert first.quantity == 5.0
        assert first.total == 2500.0

    def test_different_prices_not_merged(self):
        matches = [
            _mr(matched_position_number=1, price=500.0, quantity=1.0, total=500.0),
            _mr(matched_position_number=1, price=600.0, quantity=1.0, total=600.0),
        ]
        merged = _merge_duplicate_matches(matches)
        assert len(merged) == 2

    def test_highlight_preserved_if_any(self):
        matches = [
            _mr(matched_position_number=1, price=500.0, highlight_price=False),
            _mr(matched_position_number=1, price=500.0, highlight_price=True, highlight_reason="копейки"),
        ]
        merged = _merge_duplicate_matches(matches)
        assert len(merged) == 1
        assert merged[0].highlight_price is True

    def test_same_position_different_prices_marked_for_review(self):
        """
        Если для одной matched_position_number осталось несколько уникальных цен —
        все такие строки помечаются needs_review + highlight_price.
        """
        matches = [
            _mr(matched_position_number=1, price=500.0, quantity=1.0, total=500.0,
                needs_review=False, highlight_price=False, highlight_reason=None),
            _mr(matched_position_number=1, price=600.0, quantity=1.0, total=600.0,
                needs_review=False, highlight_price=False, highlight_reason=None),
            _mr(matched_position_number=2, price=300.0, quantity=1.0, total=300.0,
                needs_review=False, highlight_price=False, highlight_reason=None),
        ]
        merged = _merge_duplicate_matches(matches)
        # Группировка по (№, price) — все три ключа разные → 3 строки
        assert len(merged) == 3
        flagged = [m for m in merged if m.matched_position_number == 1]
        assert len(flagged) == 2
        for m in flagged:
            assert m.needs_review is True
            assert m.highlight_price is True
            assert m.highlight_reason is not None
            assert "Несколько цен" in m.highlight_reason
        # Позиция 2 — единственная цена → флаги не выставляются
        single = next(m for m in merged if m.matched_position_number == 2)
        assert single.needs_review is False
        assert single.highlight_price is False

    def test_same_position_same_price_no_extra_flag(self):
        """Обычное слияние — флаги не должны меняться."""
        matches = [
            _mr(matched_position_number=1, price=500.0, quantity=2.0, total=1000.0,
                needs_review=False, highlight_price=False),
            _mr(matched_position_number=1, price=500.0, quantity=3.0, total=1500.0,
                needs_review=False, highlight_price=False),
        ]
        merged = _merge_duplicate_matches(matches)
        assert len(merged) == 1
        assert merged[0].needs_review is False
        assert merged[0].highlight_price is False
        assert merged[0].highlight_reason is None

    def test_unmatched_positions_with_different_prices_not_flagged(self):
        """
        matched_position_number=None означает «не сопоставлено».
        Такие строки не группируются и не помечаются как «несколько цен на одну позицию».
        """
        matches = [
            _mr(matched_position_number=None, price=100.0, needs_review=True, highlight_price=False),
            _mr(matched_position_number=None, price=200.0, needs_review=True, highlight_price=False),
        ]
        merged = _merge_duplicate_matches(matches)
        assert len(merged) == 2
        # highlight_price не должен быть выставлен из-за «нескольких цен»
        for m in merged:
            assert m.highlight_price is False


class TestGenerateTemplate:
    @pytest.mark.asyncio
    async def test_basic_generation_creates_7_columns(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path)
        matches = [_mr()]
        path = await generate_filled_template(
            matches=matches,
            template_path=None,
            task_id="test_1",
            doc_number="INV-123",
            doc_date="15.03.2024",
        )
        assert Path(path).exists()
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # 7 колонок заголовков
        expected_headers = [
            "Номер документа", "Дата подписания документа", "№ позиции договора",
            "Количество (объем)", "Единица измерения",
            "Цена за единицу с НДС, руб.", "Сумма по позиции",
        ]
        for col, expected in enumerate(expected_headers, 1):
            assert ws.cell(row=1, column=col).value == expected
        # Данные в строке 2
        assert ws.cell(row=2, column=1).value == "INV-123"
        assert ws.cell(row=2, column=2).value == "15.03.2024"
        assert ws.cell(row=2, column=3).value == 1
        assert ws.cell(row=2, column=4).value == 1.0
        assert ws.cell(row=2, column=5).value == "шт."
        assert ws.cell(row=2, column=6).value == 1000.0
        assert ws.cell(row=2, column=7).value == 1000.0
        wb.close()

    @pytest.mark.asyncio
    async def test_yellow_fill_on_kopecks(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path)
        matches = [_mr(price=660.45, total=660.45)]
        path = await generate_filled_template(
            matches=matches, task_id="test_kop", doc_number="", doc_date="",
        )
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        price_cell = ws.cell(row=2, column=6)
        # Цена с копейками — должна быть подсвечена
        assert price_cell.fill.start_color.rgb in ("FFFFFF00", "00FFFF00")
        wb.close()

    @pytest.mark.asyncio
    async def test_yellow_fill_on_highlight_flag(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path)
        matches = [_mr(price=1500.0, total=1500.0, highlight_price=True)]
        path = await generate_filled_template(
            matches=matches, task_id="test_h", doc_number="", doc_date="",
        )
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        price_cell = ws.cell(row=2, column=6)
        assert price_cell.fill.start_color.rgb in ("FFFFFF00", "00FFFF00")
        wb.close()

    @pytest.mark.asyncio
    async def test_no_highlight_on_round_price(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path)
        matches = [_mr(price=1000.0, total=1000.0, highlight_price=False)]
        path = await generate_filled_template(
            matches=matches, task_id="test_noh", doc_number="", doc_date="",
        )
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        price_cell = ws.cell(row=2, column=6)
        # Fill может быть None или без жёлтого
        if price_cell.fill.start_color and price_cell.fill.start_color.rgb:
            assert price_cell.fill.start_color.rgb not in ("FFFFFF00", "00FFFF00")
        wb.close()

    @pytest.mark.asyncio
    async def test_duplicates_merged_by_default(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path)
        matches = [
            _mr(matched_position_number=1, price=500.0, quantity=2.0, total=1000.0),
            _mr(matched_position_number=1, price=500.0, quantity=3.0, total=1500.0),
        ]
        path = await generate_filled_template(
            matches=matches, task_id="test_dup", doc_number="", doc_date="",
        )
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # Должна быть одна строка данных с суммой 5.0
        assert ws.cell(row=2, column=4).value == 5.0
        assert ws.cell(row=2, column=7).value == 2500.0
        # Третья строка должна быть пустой
        assert ws.cell(row=3, column=1).value is None
        wb.close()

    @pytest.mark.asyncio
    async def test_sort_null_position_to_end(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "GENERATED_DIR", tmp_path)
        matches = [
            _mr(matched_position_number=None, price=100.0),
            _mr(matched_position_number=1, price=500.0),
        ]
        path = await generate_filled_template(
            matches=matches, task_id="test_sort", doc_number="", doc_date="",
        )
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # Первая позиция — 1, затем None
        assert ws.cell(row=2, column=3).value == 1
        assert ws.cell(row=3, column=3).value is None
        wb.close()


class TestFindFirstEmptyRow:
    def test_finds_after_header(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A", "B", "C", "D", "E", "F", "G"])
        ws.append([1, 2, 3, 4, 5, 6, 7])
        assert _find_first_empty_row(ws) == 3

    def test_empty_sheet(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        # Пустой лист — первая пустая строка 2 (после «заголовков»)
        result = _find_first_empty_row(ws)
        assert result >= 2
