"""
Generates example xlsx files for manual UI testing of the «Обработка» page.

Three scenarios are produced — each in its own subfolder:
  examples/leader_smi/   — simple UPD scenario (TV/radio ads)
  examples/veneta/        — simple UPD scenario (printer cartridges)
  examples/dlux/          — complex Act scenario (car wash, with discount + coefficient)

Run from project root:
  python examples/_generate_examples.py

Requires openpyxl (already in backend deps).
"""
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent


def save(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    wb.close()
    print(f"  + {path.relative_to(ROOT.parent)}")


# ─── 1. Лидер СМИ — простой УПД ────────────────────────────────────────


def build_leader_smi_prices() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс-лист"
    ws.append(["Прайс-лист рекламных услуг ООО «Лидер СМИ»"])
    ws.append([])
    ws.append(["№", "Наименование услуги", "Единица измерения", "Цена за единицу с НДС, руб."])
    ws.append([1, "Размещение видеоролика на телеканале «Россия 1»", "секунда", 5000])
    ws.append([2, "Размещение видеоролика на телеканале «НТВ»", "секунда", 4500])
    ws.append([3, "Размещение баннера на портале «РИА Новости»", "сутки", 12000])
    ws.append([4, "Радиоролик на радиостанции «Маяк»", "секунда", 800])
    ws.append([5, "Размещение статьи в журнале «Известия»", "полоса", 30000])
    ws.append([6, "Радиоролик на радиостанции «Эхо Москвы»", "секунда", 950])
    ws.append([7, "Размещение видеоролика на телеканале «Первый канал»", "секунда", 6500])
    return wb


def build_leader_smi_upd() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "УПД"
    ws.append(["Счет-фактура № 123/2024 от 15.03.2024"])
    ws.append(["Дата подписания: 15.03.2024"])
    ws.append(["Дата отгрузки: 14.03.2024"])
    ws.append([])
    ws.append(["№", "Наименование услуги", "Количество", "Единица измерения", "Сумма"])
    ws.append([1, "Видеоролик на телеканале Россия 1", 30, "секунда", 150000])
    ws.append([2, "Размещение баннера РИА Новости", 5, "сутки", 60000])
    ws.append([3, "Радиоролик на радиостанции Маяк", 60, "секунда", 48000])
    ws.append([4, "Видеоролик на Первом канале", 15, "секунда", 97500])
    return wb


# ─── 2. Венета (картриджи) — простой УПД ────────────────────────────────


def build_veneta_prices() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Прайс-лист расходных материалов «Венета»"])
    ws.append([])
    ws.append(["№", "Наименование товара", "Единица измерения", "Цена за единицу с НДС, руб."])
    ws.append([1, "Картридж HP CF283A", "шт", 4500])
    ws.append([2, "Картридж Canon 725", "шт", 3800])
    ws.append([3, "Картридж Brother TN-1075", "шт", 2200])
    ws.append([4, "Тонер-картридж Samsung MLT-D101S", "шт", 5100])
    ws.append([5, "Картридж Xerox 106R02773", "шт", 6300])
    ws.append([6, "Картридж Kyocera TK-1110", "шт", 2900])
    ws.append([7, "Барабан HP CF234A", "шт", 7800])
    return wb


def build_veneta_upd() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "УПД"
    ws.append(["Счет-фактура № 45/2024 от 20.03.2024"])
    ws.append(["Дата подписания: 20.03.2024"])
    ws.append([])
    ws.append(["№", "Наименование товара", "Количество", "Единица измерения", "Сумма"])
    ws.append([1, "Картридж HP CF283A", 5, "шт", 22500])
    ws.append([2, "Canon 725", 3, "шт", 11400])
    ws.append([3, "Brother TN-1075", 2, "шт", 4400])
    ws.append([4, "Samsung MLT-D101S тонер", 4, "шт", 20400])
    return wb


# ─── 3. Д-люкс (мойка) — сложный Акт со всеми нюансами ──────────────────


def build_dlux_prices() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(["Прайс-лист услуг автомойки «Д-люкс»"])
    ws.append(["Цены указаны без скидки 15%"])
    ws.append([])
    ws.append(["№", "Наименование услуги", "Единица измерения", "Цена с НДС, руб."])
    # Section header for SEDAN
    ws.append(["", "СЕДАН", "", ""])
    ws.append([1, "Мойка кузова седан", "усл. ед.", 800])
    ws.append([2, "Мойка кузова седан при сильном загрязнении", "усл. ед.", 800])
    ws.append([3, "Полировка кузова седан", "усл. ед.", 1500])
    ws.append([4, "Химчистка салона седан", "усл. ед.", 3000])
    # Section header for JEEP
    ws.append(["", "ДЖИП", "", ""])
    ws.append([5, "Мойка кузова джип", "усл. ед.", 1200])
    ws.append([6, "Мойка кузова джип при сильном загрязнении", "усл. ед.", 1200])
    ws.append([7, "Полировка кузова джип", "усл. ед.", 2200])
    ws.append([8, "Химчистка салона джип", "усл. ед.", 3500])
    # Section header for MINIBUS
    ws.append(["", "МИКРОАВТОБУС", "", ""])
    ws.append([9, "Мойка кузова микроавтобус", "усл. ед.", 1800])
    ws.append([10, "Полировка микроавтобус", "усл. ед.", 2800])
    return wb


def build_dlux_act() -> Workbook:
    """
    Тестовые кейсы в одном акте:
    - строка 1: Camry → седан, обычная мойка
    - строка 2: Camry → седан, «сильное загрязнение» (коэффициент 1.5) + полировка через «и»
    - строка 3: RAV4 → джип, мойка + химчистка через «и»
    - строка 4: Sprinter → микроавтобус, обычная мойка
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Акт"
    ws.append(["Расшифровка к акту выполненных работ № 7"])
    ws.append(["Дата подписания: 18.03.2024"])
    ws.append([])
    ws.append(["№", "Марка авто", "Наименование работы (услуги)"])
    ws.append([1, "Toyota Camry", "Мойка кузова"])
    ws.append([2, "Toyota Camry", "Мойка кузова при сильном загрязнении и полировка кузова"])
    ws.append([3, "Toyota RAV4", "Мойка кузова и химчистка салона"])
    ws.append([4, "Mercedes Sprinter", "Мойка кузова"])
    return wb


def main() -> None:
    builders = [
        ("leader_smi/01_prices_leader_smi.xlsx", build_leader_smi_prices),
        ("leader_smi/02_upd_leader_smi.xlsx", build_leader_smi_upd),
        ("veneta/01_prices_veneta.xlsx", build_veneta_prices),
        ("veneta/02_upd_veneta.xlsx", build_veneta_upd),
        ("dlux/01_prices_dlux.xlsx", build_dlux_prices),
        ("dlux/02_act_dlux.xlsx", build_dlux_act),
    ]
    print("Generating examples in", ROOT)
    for relpath, builder in builders:
        wb = builder()
        save(wb, ROOT / relpath)
    print("Done.")


if __name__ == "__main__":
    main()
